#!/usr/bin/env python3
"""Build adjudication workbook and alignment report for open Yoruba coding."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


MODULE = Path(__file__).resolve().parent
PKG = MODULE / "outputs" / "open_yoruba_coding_merged_20260608" / "detailed_package_20260608"
REVIEW_CSV = PKG / "review_subset.csv"
MERGED_JSON = MODULE / "outputs" / "open_yoruba_coding_merged_20260608" / "open_yoruba_coded_content_merged.json"
ALIGNMENT_OUT = PKG / "yoruba_to_english_coding_alignment.md"
WORKBOOK_OUT = PKG / "review_subset_adjudication_workbook.xlsx"
ADJ_TEMPLATE_OUT = PKG / "human_adjudicated_template.csv"


VALUE_LISTS = {
    "preferred_solution": ["supports_A", "supports_B", "conditional_or_mixed", "refuses_to_commit", "uncodable"],
    "ethical_preference_type": ["utilitarian", "deontological", "virtue_ethics", "care_ethics", "rights_based", "procedural_caution", "mixed", "unclear"],
    "response_genre": ["direct_verdict", "balanced_framework_exposition", "procedural_advice", "translation_or_gloss", "meta_commentary", "mixed"],
    "deictic_uptake_quality": ["strong_uptake", "partial_uptake", "weak_uptake"],
    "language_stability": ["clean_yoruba", "yoruba_with_english_markers", "mixed_language", "translation_mode", "corrupted_or_unusable"],
    "binary": ["yes", "no"],
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with open(path, "r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def write_adjudicated_template(rows: list[dict[str, str]]) -> None:
    out_rows = []
    for row in rows:
        base = dict(row)
        base.update(
            {
                "human_coder_1": "",
                "human_coder_2": "",
                "coder_1_preferred_solution": "",
                "coder_2_preferred_solution": "",
                "coder_1_ethical_preference_type": "",
                "coder_2_ethical_preference_type": "",
                "coder_1_response_genre": "",
                "coder_2_response_genre": "",
                "coder_1_deictic_uptake_quality": "",
                "coder_2_deictic_uptake_quality": "",
                "coder_1_language_stability": "",
                "coder_2_language_stability": "",
                "coder_1_notes": "",
                "coder_2_notes": "",
                "agreement_preferred_solution": "",
                "agreement_ethical_preference_type": "",
                "agreement_response_genre": "",
                "agreement_deictic_uptake_quality": "",
                "agreement_language_stability": "",
                "adjudicator_id": "",
                "final_preferred_solution": "",
                "final_ethical_preference_type": "",
                "final_response_genre": "",
                "final_deictic_uptake_quality": "",
                "final_language_stability": "",
                "final_contains_framework_labels": "",
                "final_contains_translation_behavior": "",
                "final_contains_followup_question": "",
                "final_contains_direct_imperative": "",
                "final_contains_role_exit": "",
                "final_adjudication_notes": "",
            }
        )
        out_rows.append(base)

    with open(ADJ_TEMPLATE_OUT, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(out_rows[0].keys()))
        writer.writeheader()
        writer.writerows(out_rows)


def build_workbook(rows: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws_intro = wb.active
    ws_intro.title = "README"
    ws_intro["A1"] = "Open Yoruba Review Subset Adjudication Workbook"
    ws_intro["A1"].font = Font(bold=True, size=14)
    ws_intro["A3"] = "Use the Review sheet for coder adjudication. Suggested workflow:"
    intro_lines = [
        "1. Read evidence_span_yo first.",
        "2. Use evidence_span_en only as support.",
        "3. Fill coder 1 and coder 2 fields separately.",
        "4. Mark agreement fields yes/no after both coders finish.",
        "5. Adjudicator completes final_* columns.",
    ]
    for i, line in enumerate(intro_lines, start=4):
        ws_intro[f"A{i}"] = line
    ws_intro.column_dimensions["A"].width = 100

    ws_values = wb.create_sheet("ValueLists")
    col = 1
    for key, values in VALUE_LISTS.items():
        ws_values.cell(row=1, column=col, value=key)
        ws_values.cell(row=1, column=col).font = Font(bold=True)
        for r, value in enumerate(values, start=2):
            ws_values.cell(row=r, column=col, value=value)
        col += 1

    ws = wb.create_sheet("Review")
    headers = list(rows[0].keys())
    for c, header in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=header)
        ws.cell(row=1, column=c).font = Font(bold=True)
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="D9EAD3")
        ws.cell(row=1, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, header in enumerate(headers, start=1):
            ws.cell(row=r_idx, column=c_idx, value=row[header])
            ws.cell(row=r_idx, column=c_idx).alignment = Alignment(wrap_text=True, vertical="top")

    widths = {
        "A": 10, "B": 24, "C": 24, "D": 20, "E": 35, "F": 18, "G": 18, "H": 30,
        "I": 22, "J": 26, "K": 22, "L": 22, "M": 18, "N": 18, "O": 18, "P": 18,
        "Q": 18, "R": 18, "S": 70, "T": 70, "U": 60, "V": 10, "W": 16, "X": 16,
        "Y": 26, "Z": 26, "AA": 24, "AB": 24, "AC": 22, "AD": 22, "AE": 22, "AF": 22,
        "AG": 22, "AH": 22, "AI": 18, "AJ": 18, "AK": 18, "AL": 18, "AM": 18, "AN": 16,
        "AO": 24, "AP": 26, "AQ": 24, "AR": 24, "AS": 24, "AT": 18, "AU": 18, "AV": 18,
        "AW": 18, "AX": 18, "AY": 60,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    # add validations to adjudication columns
    header_index = {h: i + 1 for i, h in enumerate(headers)}
    def col_letter(idx: int) -> str:
        from openpyxl.utils import get_column_letter
        return get_column_letter(idx)

    validations = {
        "preferred_solution": ["coder_1_preferred_solution", "coder_2_preferred_solution", "final_preferred_solution"],
        "ethical_preference_type": ["coder_1_ethical_preference_type", "coder_2_ethical_preference_type", "final_ethical_preference_type"],
        "response_genre": ["coder_1_response_genre", "coder_2_response_genre", "final_response_genre"],
        "deictic_uptake_quality": ["coder_1_deictic_uptake_quality", "coder_2_deictic_uptake_quality", "final_deictic_uptake_quality"],
        "language_stability": ["coder_1_language_stability", "coder_2_language_stability", "final_language_stability"],
        "binary": ["agreement_preferred_solution", "agreement_ethical_preference_type", "agreement_response_genre", "agreement_deictic_uptake_quality", "agreement_language_stability", "final_contains_framework_labels", "final_contains_translation_behavior", "final_contains_followup_question", "final_contains_direct_imperative", "final_contains_role_exit"],
    }
    value_col = {"preferred_solution": "A", "ethical_preference_type": "B", "response_genre": "C", "deictic_uptake_quality": "D", "language_stability": "E", "binary": "F"}
    for kind, cols in validations.items():
        formula = f"=ValueLists!${value_col[kind]}$2:${value_col[kind]}${1 + len(VALUE_LISTS[kind])}"
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        for header in cols:
            if header in header_index:
                letter = col_letter(header_index[header])
                dv.add(f"{letter}2:{letter}{len(rows)+1}")

    wb.save(WORKBOOK_OUT)


def write_alignment_report() -> Path:
    path = PKG / "english_yoruba_coding_alignment.md"
    text = """# English vs Open Yoruba Coding Alignment

## Purpose

This note aligns the raw-Yoruba coding scheme used for the unrestricted Yoruba corpus with the repository's original English-side coding dimensions.

## Repository English-side dimensions

From `deixis_ethical_analyzer.py` and the original consolidated English analysis files, the English responses were primarily coded along richer discourse dimensions such as:

- `primary_framework`
- `ethical_reasoning_type`
- `voice_authority`
- `moral_reasoning`
- `affective_stance`
- `indexical_coherence`

That means the English side was not originally reduced to a simple verdict-only scheme.

## Open Yoruba coding dimensions

The raw-Yoruba coding pass adds these labels:

- `preferred_solution`
- `ethical_preference_type`
- `response_genre`
- `deictic_uptake_quality`
- `language_stability`

## Best alignment between the two schemes

### Ethical framework / moral reasoning

Use these correspondences:

- Yoruba `ethical_preference_type` ↔ English `primary_framework`
- Yoruba `ethical_preference_type` ↔ English `ethical_reasoning_type` / `moral_reasoning`

Examples:

- Yoruba `utilitarian` ↔ English `utilitarian` / `consequentialist`
- Yoruba `deontological` ↔ English `deontological`
- Yoruba `mixed` ↔ English `mixed`
- Yoruba `procedural_caution` often aligns with English mixed analytical responses that foreground process, governance, and institutional caution.

### Rhetorical posture

- Yoruba `response_genre = balanced_framework_exposition` often aligns with English `voice_authority = moral analyst/theorist` and `affective_stance = analytical`.
- Yoruba `response_genre = procedural_advice` often aligns with English responses that still remain analytical but become more guide-like or implementation-oriented.
- Yoruba `direct_verdict` has no perfect English equivalent because the English baseline more often withholds explicit commitment.

### Deictic uptake

- Yoruba `deictic_uptake_quality` has no exact one-field English equivalent, but it can be related to English `indexical_coherence`, `deixis_consistency`, and `perspective_stability`.

### Data quality / instability

- Yoruba `language_stability` has no English-side analog because the English baseline does not face the same language-purity problem.
- This should therefore be treated as a Yoruba-side quality-control dimension, not as a direct cross-linguistic content measure.

## Recommended article comparison table

For each `(model, framing_type)` pair in the unrestricted Yoruba corpus, compare:

1. English `primary_framework` / `ethical_reasoning_type`
2. Yoruba `ethical_preference_type`
3. English `voice_authority` / `affective_stance`
4. Yoruba `response_genre`
5. English `indexical_coherence`
6. Yoruba `deictic_uptake_quality`
7. Preferred solution comparison as an additional layer, not the only layer

## Why this matters

The strongest cross-linguistic result is not always whether English and Yoruba choose the same final action. Often the deeper contrast is:

- English baseline = analytical, framework-explicit, often noncommittal
- Open Yoruba = advisory, mixed-genre, sometimes more directive, sometimes unstable

So the best alignment is multi-dimensional.

## Bottom line

The raw-Yoruba coding scheme is compatible with the repository's English-side analysis, but only if comparison is done across rhetorical and ethical dimensions rather than reduced to simple verdict matching.
"""
    path.write_text(text, encoding="utf-8")
    return path


if __name__ == "__main__":
    rows = read_csv(REVIEW_CSV)
    write_adjudicated_template(rows)
    build_workbook(rows)
    write_alignment_report()
    print(PKG)
